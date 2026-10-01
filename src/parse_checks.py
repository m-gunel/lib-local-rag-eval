"""Быстрые проверки парсинга на корпусе стенда: после каждой правки парсеров, без демона.

Запуск интерпретатором демона (парсеры проекта, pymupdf, python-pptx, openpyxl):
  PYTHONDONTWRITEBYTECODE=1 daemon-venv/bin/python src/parse_checks.py check --view --label R7
      базовая линия: чанки из work/indexer_view_corpus (как в R7), без метаданных
  LIB_LOCAL_RAG=<проект> PYTHONDONTWRITEBYTECODE=1 daemon-venv/bin/python src/parse_checks.py check --live --label <метка>
      чанки и метаданные текущими парсерами проекта через indexer_view.records_of; кэш
      work/parse_cache/<ключ формата>/<sha1 файла>.json, отчёт work/parse_checks/<метка>.json
  daemon-venv/bin/python src/parse_checks.py delta work/parse_checks/R7.json work/parse_checks/<метка>.json
      какие числа сдвинулись
  daemon-venv/bin/python src/parse_checks.py compare --run P1 --base R7
      парное сравнение прогонов по срезам: формат, набор, сценарий
  daemon-venv/bin/python src/parse_checks.py freeze-xlsx-gold
      один раз, ДО правки парсера xlsx: эталон ответов xlsx-запросов в data/xlsx_gold.jsonl

Ключ кэша = sha1(модуль парсера формата + все прочие файлы src/parsers + config.py + doc_meta.py +
r7_office.py + файлы токенизатора + настройки сплиттера и excel + версии библиотек разбора).
Правка pdf.py меняет ключ только для pdf; правка общего модуля — для всех форматов. Чистить кэш
руками не нужно, прежние версии чанков остаются рядом.
"""

import argparse
import collections
import datetime
import functools
import hashlib
import importlib.util
import json
import math
import os
import re
import statistics
import sys
import time
import types
import unicodedata
from pathlib import Path

EVAL = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(EVAL / "src"))  # metrics.py, report.py стенда
DATA, WORK = EVAL / "data", EVAL / "work"
VIEW, CACHE, OUT, TR = WORK / "indexer_view_corpus", WORK / "parse_cache", WORK / "parse_checks", WORK / "translate"
PROJECT = Path(os.environ.get("LIB_LOCAL_RAG", "/Users/gunel30/Downloads/pip_local_rag_0110"))
XLSX_GOLD = DATA / "xlsx_gold.jsonl"

TOK = re.compile(r"[0-9a-zа-я]+")
HYPH_DIAG = re.compile("[а-яёa-z][\u00ad\u2011\\-\u2010]\\s*\\n\\s*[а-яёa-z]", re.I)  # регулярка постановки
HYPH_CTRL = re.compile("[а-яёa-z][\x16-\x1c]\\s*\\n\\s*[а-яёa-z]", re.I)  # дефис, извлечённый управляющим символом
CTRL = re.compile("[\x00-\x08\x0b\x0c\x0e-\x1f\x7f-\x9f]")
MONTHS = "января февраля марта апреля мая июня июля августа сентября октября ноября декабря".split()
BANK = {  # строгое имя (как в постановке) и с псевдонимами
    "alfa": (r"альфа-банк", r"альфа[\s\-\u2011–]*банк"), "domrf": (r"дом\.рф", r"дом[\s.]*рф"),
    "gpb": (r"газпромбанк", r"газпромбанк"), "mkb": (r"московский кредитный банк", r"московск\w*\s+кредитн\w*\s+банк|\bмкб\b"),
    "rshb": (r"россельхозбанк", r"россельхозбанк|сельскохозяйственн\w*\s+банк"), "sber": (r"сбербанк", r"сбербанк|(?<![а-я])сбер\b"),
    "sovkom": (r"совкомбанк", r"совкомбанк"), "tbank": (r"т-банк", r"(?<![а-яa-z])т[\-\u2011–]?\s?банк|тинькофф"), "vtb": (r"втб", r"\bвтб\b"),
}
# Строка-префикс, которую конвейер проекта ставит сверху чанка («Документ: …»): для проверок
# собственного содержимого чанка её снимаем.
PREFIX_LINE = re.compile(r"^(Документ|Раздел|Лист|Таблица): [^\n]*\n")


def norm(s: str) -> str:
    return re.sub(r"\s+", " ", str(s).lower().replace("ё", "е")).strip()


def tok(s: str) -> list[str]:
    return TOK.findall(s.lower().replace("ё", "е"))


def grams(t: list[str], n: int = 3) -> set:
    return {tuple(t[i:i + n]) for i in range(len(t) - n + 1)} if len(t) >= n else ({tuple(t)} if t else set())


def sha1_file(p: Path) -> str:
    h = hashlib.sha1()
    with open(p, "rb") as f:
        for b in iter(lambda: f.read(1 << 20), b""):
            h.update(b)
    return h.hexdigest()


def jl(p: Path) -> list[dict]:
    return [json.loads(line) for line in open(p, encoding="utf-8") if line.strip()]


def own_text(chunk_text: str) -> str:
    """Текст чанка без строк-префиксов конвейера."""
    t = chunk_text
    while True:
        m = PREFIX_LINE.match(t)
        if not m:
            return t
        t = t[m.end():]


MANIFEST = jl(DATA / "manifest.jsonl")
BY_KEY = {m["key"]: m for m in MANIFEST}


# ---------- источник чанков ----------

class ViewSource:
    """Чанки, по которым оценивался R7 (work/indexer_view_corpus); метаданных там нет."""
    live = False

    def chunks(self, m):
        d = json.load(open(VIEW / (m["doc_id"].replace("/", "__") + ".chunks.json"), encoding="utf-8"))
        return [{"text": t} for t in d["chunks"]], None, d.get("error"), None


class LiveSource:
    """Чанки и метаданные текущими парсерами проекта, кэш по ключу версии парсера."""
    live = True

    def __init__(self):
        os.environ.setdefault("INDEXER_VIEW_HOME", str(WORK / "indexer_view_home"))
        import indexer_view  # импорт парсеров проекта и set_config()
        from src.parsers.base import file_extension, get_parser, get_parsers_registry
        from src.utils import get_config
        self.iv, self.cfg, self.get_parser, self.ext = indexer_view, get_config(), get_parser, file_extension
        self.format_modules = {Path(sys.modules[c.__module__].__file__).resolve() for c in get_parsers_registry().values()}
        self.keys: dict[str, str] = {}

    def parser_key(self, ext: str) -> str:
        if ext not in self.keys:
            from importlib.metadata import version
            proj = self.iv.PROJECT
            own = Path(sys.modules[type(self.get_parser(ext)).__module__].__file__).resolve()
            common = [f for f in sorted((proj / "src/parsers").rglob("*"))
                      if f.is_file() and "__pycache__" not in f.parts and f.resolve() not in self.format_modules]
            files = [own, *common, proj / "src/config.py", proj / "src/common/doc_meta.py", proj / "src/utils/r7_office.py"]
            mp = Path(self.cfg.embed_model_cfg.model_path)
            mp = mp if mp.is_absolute() else proj / mp
            files += [mp / "openvino_tokenizer.xml", mp / "openvino_tokenizer.bin"]
            h = hashlib.sha1()
            for f in files:
                if f.exists():
                    h.update(str(f.relative_to(proj) if f.is_relative_to(proj) else f.name).encode())
                    h.update(f.read_bytes())
            libs = {}
            for n in ("pymupdf", "python-pptx", "openpyxl", "docx2txt", "langchain-text-splitters", "openvino-tokenizers", "lxml"):
                try:
                    libs[n] = version(n)
                except Exception:
                    libs[n] = None
            cfg = {"splitter": self.cfg.embed_splitter_cfg.model_dump(), "excel": self.cfg.excel_config.model_dump(), "libs": libs}
            h.update(json.dumps(cfg, sort_keys=True, default=str).encode())
            self.keys[ext] = h.hexdigest()[:16]
        return self.keys[ext]

    def chunks(self, m):
        path = Path(m["path"])
        c = CACHE / self.parser_key(self.ext(path)) / (sha1_file(path) + ".json")
        if c.exists():
            d = json.load(open(c, encoding="utf-8"))
            return d["chunks"], d.get("meta"), d["error"], d["seconds"]
        t = time.perf_counter()
        try:
            (ch, meta), err = self.iv.records_of(path), None
        except Exception as e:  # битый файл — тоже результат
            ch, meta, err = [], None, f"{type(e).__name__}: {e}"
        sec = round(time.perf_counter() - t, 3)
        c.parent.mkdir(parents=True, exist_ok=True)
        c.write_text(json.dumps({"file": str(path), "error": err, "seconds": sec, "meta": meta, "chunks": ch},
                                ensure_ascii=False), encoding="utf-8")
        return ch, meta, err, sec


def token_counter(src):
    """count_tokens модели (live) или его копия WordPiece на чистом Python (расхождение с OpenVINO 0,015%)."""
    if src.live:
        return src.cfg.ov_model.count_tokens
    tj = json.load(open(PROJECT / "models/rubert-tiny2_002/tokenizer.json", encoding="utf-8"))
    vocab, cache = tj["model"]["vocab"], {}

    def wp(w):
        if w not in cache:
            n, s = 0, 0
            if len(w) > 100:
                cache[w] = 1
                return 1
            while s < len(w):
                e = len(w)
                while s < e and (w[s:e] if s == 0 else "##" + w[s:e]) not in vocab:
                    e -= 1
                if s == e:
                    n, s = 1, len(w)  # [UNK] на всё слово
                    break
                n, s = n + 1, e
            cache[w] = n
        return cache[w]

    def punct(ch):
        cp = ord(ch)
        return 33 <= cp <= 47 or 58 <= cp <= 64 or 91 <= cp <= 96 or 123 <= cp <= 126 or unicodedata.category(ch).startswith("P")

    def count(t):
        clean = "".join(" " if ch in "\t\n\r" or unicodedata.category(ch) == "Zs" else ch for ch in t
                        if ord(ch) not in (0, 0xFFFD) and unicodedata.category(ch) not in ("Cc", "Cf") or ch in "\t\n\r")
        n = 0
        for w in clean.split():
            buf = ""
            for ch in w:
                if punct(ch):
                    n += (wp(buf) if buf else 0) + 1
                    buf = ""
                else:
                    buf += ch
            n += wp(buf) if buf else 0
        return min(n + 2, 2048)

    return count


def load_all(src) -> dict[str, dict]:
    docs = {}
    for m in MANIFEST:
        ch, meta, err, sec = src.chunks(m)
        docs[m["doc_id"]] = {"m": m, "chunks": [c["text"] for c in ch], "records": ch, "meta": meta,
                             "error": err, "seconds": sec}
    return docs


def pct(a, b):
    return round(100 * a / b, 1) if b else None


# ---------- проверки ----------

def check_same_as_view(docs):
    """Совпадают ли чанки (текст) с теми, по которым оценивался R7."""
    same, diff = 0, []
    for doc_id, d in docs.items():
        v = json.load(open(VIEW / (doc_id.replace("/", "__") + ".chunks.json"), encoding="utf-8"))["chunks"]
        if v == d["chunks"]:
            same += 1
        else:
            diff.append(doc_id)
    return {"docs": len(docs), "same_as_R7": same, "differ_examples": diff[:10]}


def check_volume(docs, count):
    out = collections.defaultdict(lambda: collections.Counter())
    for d in docs.values():
        m = d["m"]
        c = out[f"{m['dataset']}/{m['format']}"]
        c["docs"] += 1
        c["docs_without_chunks"] += not d["chunks"]
        c["errors"] += bool(d["error"])
        c["chunks"] += len(d["chunks"])
        c["chars"] += sum(map(len, d["chunks"]))
        toks = [count(x) for x in d["chunks"]]
        c["tokens"] += sum(toks)
        c["chunks_over_560_tokens"] += sum(t > 560 for t in toks)  # 500 + префикс «Документ: …»
        if d["seconds"] is not None:
            c["parse_ms"] += int(1000 * d["seconds"])
    tot = collections.Counter()
    for c in out.values():
        tot.update(c)
    out["ИТОГО"] = tot
    return {k: dict(v) for k, v in sorted(out.items())}


def check_pdf_text(docs):
    res = {}
    for ds in sorted({d["m"]["dataset"] for d in docs.values() if d["m"]["format"] == "pdf"}):
        c = collections.Counter()
        for d in docs.values():
            if d["m"]["dataset"] != ds or d["m"]["format"] != "pdf":
                continue
            ch = [own_text(x) for x in d["chunks"]]
            line_cnt = collections.Counter()
            for x in ch:
                for line in {norm(s) for s in x.split("\n")}:
                    if 5 <= len(line) <= 80 and not line.isdigit():
                        line_cnt[line] += 1
            boiler = {s for s, n in line_cnt.items() if n >= max(3, 0.1 * len(ch))}  # строка в ≥10% чанков: колонтитул
            for x in ch:
                c["chunks"] += 1
                c["breaks"] += len(HYPH_DIAG.findall(x))
                c["breaks_ctrl"] += len(HYPH_CTRL.findall(x))
                c["chunks_soft_hyphen"] += "\u00ad" in x
                c["chunks_ctrl_chars"] += bool(CTRL.search(x))
                c["chunks_vestnik_caps"] += bool(re.search(r"ВЕСТНИК\s+БАНКА\s+РОССИИ", x))
                c["chunks_vestnik_any_case"] += bool(re.search(r"вестник\s+банка\s+россии", x, re.I))
                c["chunks_repeated_line"] += any(norm(s) in boiler for s in x.split("\n"))
        n = c["chunks"]
        res[ds] = {"chunks": n, "breaks_per_chunk": round(c["breaks"] / n, 3) if n else None, "breaks_total": c["breaks"],
                   "ctrl_breaks_total": c["breaks_ctrl"], **{k + "_%": pct(v, n) for k, v in c.items() if k.startswith("chunks_")}}
    return res


def check_context(docs):
    """Контекст документа в чанке: название/ключ документа внутри текста чанка."""
    res = {}
    by_ds = collections.defaultdict(list)
    for d in docs.values():
        by_ds[d["m"]["dataset"]].append(d)
    for ds, dd in sorted(by_ds.items()):
        c = collections.Counter()
        for d in dd:
            m = d["m"]
            words = [w[:5] for w in tok(re.sub(r"\.(pdf|docx|pptx|txt|xlsx)$", "", m["title"])) if len(w) >= 4 and not w.isdigit()][:6]
            key = norm(m.get("key", ""))
            for x in d["chunks"]:
                t = norm(x)
                c["n"] += 1
                c["title_words_50%"] += bool(words) and sum(w in t for w in words) / len(words) >= 0.5
                if ds == "cbr_acts":  # номер акта, «154-И»: дефис любой, пробелы любые
                    num, _, suf = key.partition("-")
                    c["act_number"] += bool(re.search(rf"(?<!\d){re.escape(num)}\s*\W?\s*{re.escape(suf)}(?![а-я])", t))
                if ds == "ai360":
                    bank, year = m["key"].split("_")[:2]
                    strict, alias = BANK[bank]
                    c["bank_strict"] += bool(re.search(strict, t))
                    c["bank_alias"] += bool(re.search(alias, t))
                    c["year"] += bool(re.search(rf"(?<!\d){year}(?!\d)", t))
                    c["bank_alias_and_year"] += bool(re.search(alias, t)) and bool(re.search(rf"(?<!\d){year}(?!\d)", t))
        res[ds] = {"chunks": c["n"], **{k + "_%": pct(v, c["n"]) for k, v in c.items() if k != "n"}}
    return res


def pptx_full_text(path) -> tuple[str, collections.Counter]:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    st, out = collections.Counter(), []

    def walk(shapes):
        for sh in shapes:
            if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
                st["groups"] += 1
                walk(sh.shapes)
            elif getattr(sh, "has_table", False) and sh.has_table:
                st["tables"] += 1
                out.extend(" | ".join(c.text for c in r.cells) for r in sh.table.rows)
            elif getattr(sh, "has_text_frame", False) and sh.has_text_frame and sh.text_frame.text:
                out.append(sh.text_frame.text)

    for s in Presentation(path).slides:
        walk(s.shapes)
        if s.has_notes_slide and s.notes_slide.notes_text_frame is not None and s.notes_slide.notes_text_frame.text:
            st["notes"] += 1
            out.append(s.notes_slide.notes_text_frame.text)
    return "\n".join(out), st


def check_pptx(docs):
    rows, st = [], collections.Counter()
    for d in docs.values():
        m = d["m"]
        if m["format"] != "pptx":
            continue
        full, s = pptx_full_text(m["path"])
        st.update(s)
        ref = grams(tok(full))
        cov = len(ref & grams(tok("\n".join(d["chunks"])))) / len(ref) if ref else 1.0
        rows.append((round(cov, 3), m["title"][:60]))
    rows.sort()
    covs = [r[0] for r in rows]
    return {"files": len(rows), "mean_cov_vs_full_python_pptx": round(statistics.mean(covs), 3),
            "files_cov_below_0.8": sum(c < 0.8 for c in covs), "tables": st["tables"], "groups": st["groups"],
            "notes": st["notes"], "worst": rows[:5]}


MD_SEP = re.compile(r"^\s*\|?\s*:?-{3,}")


def md_lines(p: Path) -> list[str]:
    return [line for line in p.read_text(encoding="utf-8").splitlines() if line.strip() and not MD_SEP.match(line)]


def md_clean(s: str) -> str:
    return re.sub(r"[#*_`|>\[\]]", " ", s)


def check_source_md(docs):
    """ZX Bank и ADGM: доля 3-грамм исходного русского md (из него собраны docx/pdf/pptx) в чанках."""
    res = collections.defaultdict(list)
    for d in docs.values():
        m = d["m"]
        f = TR / "docs" / (m.get("key", "") + ".ru.md")
        if not f.exists():
            continue
        ref = grams(tok(md_clean(f.read_text(encoding="utf-8"))))
        res[f"{m['dataset']}/{m['format']}"].append(len(ref & grams(tok("\n".join(d["chunks"])))) / len(ref))
    return {k: {"docs": len(v), "mean": round(statistics.mean(v), 3), "min": round(min(v), 3)} for k, v in sorted(res.items())}


# ---------- «ответ есть в индексе» ----------

def best_window(fact: str, lines: list[str], maxlen: int = 8):
    ft = collections.Counter(tok(fact))
    n, best = sum(ft.values()), (0.0, None)
    for i in range(len(lines)):
        wc = collections.Counter()
        for j in range(i, min(len(lines), i + maxlen)):
            wc += collections.Counter(tok(lines[j]))
            inter = sum((ft & wc).values())
            rec, prec = inter / n, inter / max(1, sum(wc.values()))
            f1 = 2 * rec * prec / max(1e-9, rec + prec)
            if f1 > best[0]:
                best = (f1, (i, j))
    return best


def reference_passages():
    """Эталонные фрагменты на русском: (набор, qid, doc_id, текст).

    ZX Bank и ObliQA: факты даны по-английски (work/translate/queries_*.jsonl, facts_en);
    находим их строки в исходном en.md и берём те же строки ru.md — у всех 110 пар en/ru
    совпадает построчная структура. Факты ZX, пересказанные своими словами (F1 < 0.8), пропускаем.
    FAQ ЦБ: ответ из raw/cbr_faq/pairs.jsonl, если он целиком есть на странице (answer_on_page = 1).
    """
    out, skipped = [], collections.Counter()
    for fn, ds in (("queries_zx.jsonl", "zx"), ("queries_obliqa.jsonl", "obliqa")):
        for q in jl(TR / fn):
            for f in q["facts_en"]:
                u = f["unit_id"]
                en, ru = md_lines(TR / "docs" / f"{u}.en.md"), md_lines(TR / "docs" / f"{u}.ru.md")
                f1, span = best_window(f["text"], en)
                if f1 < 0.8:
                    skipped[ds] += 1
                    continue
                out.append((ds, q["qid"], BY_KEY[u]["doc_id"], md_clean("\n".join(ru[span[0]:span[1] + 1]))))
    queries = {q["qid"]: q for q in jl(DATA / "queries.jsonl")}
    pairs = {p["question"].strip(): p for p in jl(EVAL / "raw/cbr_faq/pairs.jsonl")}
    gold = collections.defaultdict(dict)
    for line in open(DATA / "qrels.tsv", encoding="utf-8").read().splitlines()[1:]:
        q, d, g, _ = line.split("\t")
        gold[q][d] = int(g)
    for qid, q in queries.items():
        if q["dataset"] == "cbr_faq":
            p = pairs.get(q["original"].strip())
            if p and p.get("answer_on_page") == 1.0:
                out += [("cbr_faq", qid, d, p["answer"]) for d, g in gold[qid].items() if g == 3]
    return out, dict(skipped)


def check_answers(docs, threshold=0.8):
    refs, skipped = reference_passages()
    res = collections.defaultdict(list)
    gcache = {}
    for ds, qid, doc_id, text in refs:
        if doc_id not in gcache:
            ch = docs[doc_id]["chunks"]
            gcache[doc_id] = ([grams(tok(x)) for x in ch], grams(tok("\n".join(ch))))
        per_chunk, whole = gcache[doc_id]
        ref = grams(tok(text))
        if not ref:
            continue
        cd = len(ref & whole) / len(ref)
        cc = max((len(ref & g) / len(ref) for g in per_chunk), default=0.0)
        res[f"{ds}/{docs[doc_id]['m']['format']}"].append((cd, cc))
    out = {k: {"refs": len(v), "found_in_doc_%": pct(sum(a >= threshold for a, _ in v), len(v)),
               "found_in_one_chunk_%": pct(sum(b >= threshold for _, b in v), len(v)),
               "mean_cov_doc": round(statistics.mean(a for a, _ in v), 3)} for k, v in sorted(res.items())}
    out["_skipped_non_extractive"] = skipped
    out["ai360/pdf numbers"] = check_ai360_numbers(docs)
    return out


def check_ai360_numbers(docs):
    """ai360: числа из gold_answer (≥4 цифр, не год) в тексте эталонного отчёта. Сторожевая метрика:
    часть чисел — вычисленные (доли, разности), их в отчёте нет и не будет."""
    qids = {q["qid"] for q in jl(DATA / "queries.jsonl")}
    num = re.compile(r"\d[\d\s ,.]*\d")

    def flat(s):
        return re.sub(r"(?<=\d)[\s ,.](?=\d)", "", s)

    texts, c = {}, collections.Counter()
    for r in jl(EVAL / "raw/ai360/data/dataset.jsonl"):
        if "ai_" + r["question_id"][2:] not in qids:
            continue
        keys = [e["doc_id"] for e in r["gold_evidence"] if e["doc_id"] in BY_KEY]
        nums = [n for n in (re.sub(r"[\s ,.]", "", x) for x in num.findall(r["gold_answer"])) if len(n) >= 4 and not re.fullmatch(r"20[12]\d", n)]
        if not keys or not nums:
            continue
        for k in keys:
            texts.setdefault(k, flat("\n".join(docs[BY_KEY[k]["doc_id"]]["chunks"])))
        t = "".join(texts[k] for k in keys)
        found = sum(n in t for n in nums)
        c["questions"] += 1
        c["numbers"] += len(nums)
        c["numbers_found"] += found
        c["all_found"] += found == len(nums)
    return {"questions": c["questions"], "numbers_found_%": pct(c["numbers_found"], c["numbers"]),
            "questions_all_numbers_found_%": pct(c["all_found"], c["questions"])}


# ---------- гигиена текста, короткие чанки, повторы ----------

HYGIENE = {
    "nbsp": "[\u00a0\u202f\u2007]",
    "zero_width": "[\u200b\u200c\u200d\u2060]",
    "bom": "\ufeff",
    "soft_hyphen": "\u00ad",
    "ctrl": "[\x00-\x08\x0b\x0c\x0e-\x1f\x7f-\x9f]",
    "symbol_pua": "[\uf000-\uf0ff]",
    "cr": "\r",
    "tab": "\t",
    "spaces_2plus": "(?<=\\S)  +(?=\\S)",
    "trailing_space": " +\n",
    "newlines_3plus": "\n{3,}",
}
EMOJI = re.compile("[\U0001F000-\U0001FAFF☀-➿⬀-⯿]")


def check_hygiene(docs):
    res = {}
    by_fmt = collections.defaultdict(list)
    for d in docs.values():
        by_fmt[d["m"]["format"]].append(d)
    for fmt, dd in sorted(by_fmt.items()):
        c, n = collections.Counter(), 0
        for d in dd:
            for x in d["chunks"]:
                n += 1
                for k, p in HYGIENE.items():
                    c[k] += bool(re.search(p, x))
        res[fmt] = {"chunks": n, **{k + "_%": pct(c[k], n) for k in HYGIENE}}
    return res


def check_short_dups(docs):
    """Короткие чанки (меньше 3 слов из букв, без строк-префиксов) и точные повторы внутри документа."""
    res = {}
    by_fmt = collections.defaultdict(list)
    for d in docs.values():
        by_fmt[d["m"]["format"]].append(d)
    words = re.compile(r"[A-Za-zА-Яа-яЁё]{2,}")
    for fmt, dd in sorted(by_fmt.items()):
        c = collections.Counter()
        for d in dd:
            seen = set()
            for x in d["chunks"]:
                body = own_text(x)
                c["chunks"] += 1
                c["short"] += len(words.findall(body)) < 3
                key = norm(body)
                c["dup_in_doc"] += key in seen
                seen.add(key)
        res[fmt] = {"chunks": c["chunks"], "short_%": pct(c["short"], c["chunks"]), "short": c["short"],
                    "dup_in_doc": c["dup_in_doc"]}
    return res


# ---------- метаданные ----------

# Поля как в src/common/doc_meta.py проекта.
DOC_FIELDS = ("title", "author", "last_modified_by", "subject", "doc_created", "doc_modified", "page_count",
              "language", "mime", "size", "sha256", "date_created", "date_modified", "source_uri",
              "origin_url", "encoding")
CHUNK_FIELDS = ("page_start", "line_start", "section", "kind", "table_id", "row_start", "char_start")
TITLE_JUNK = re.compile(r"^(microsoft (word|excel|powerpoint)\s*-|please read first|документ\d*$|книга\d*$|"
                        r"презентация\d*$|untitled|без названия|[a-z]:\\|/)", re.I)


def filled(v) -> bool:
    return v not in (None, "", [], {})


def check_metadata(docs):
    """Полнота и качество метаданных по форматам (только для live: в view метаданных нет)."""
    res = {}
    by_fmt = collections.defaultdict(list)
    for d in docs.values():
        by_fmt[d["m"]["format"]].append(d)
    for fmt, dd in sorted(by_fmt.items()):
        metas = [d["meta"] for d in dd if d["meta"]]
        if not metas:
            res[fmt] = {"docs": len(dd), "docs_with_meta": 0}
            continue
        doc_fill = {f: pct(sum(filled(m.get(f)) for m in metas), len(metas)) for f in DOC_FIELDS}
        c = collections.Counter()
        for m in metas:
            dc, dm = m.get("date_created"), m.get("date_modified")
            c["created_after_modified"] += bool(dc and dm and dc > dm)
            c["title_junk"] += bool(m.get("title") and TITLE_JUNK.search(str(m["title"]).strip()))
            c["meta_with_emoji"] += any(isinstance(v, str) and EMOJI.search(v) for v in m.values())
            c["meta_with_ctrl"] += any(isinstance(v, str) and CTRL.search(v) for v in m.values())
        recs = [r for d in dd for r in d["records"]]
        chunk_fill = {f: pct(sum(filled(r.get(f)) for r in recs), len(recs)) for f in CHUNK_FIELDS}
        tables = [r for r in recs if r.get("kind") in ("table", "table_row")]
        res[fmt] = {"docs": len(dd), "docs_with_meta": len(metas), "doc_fields_filled_%": doc_fill,
                    "chunks": len(recs), "chunk_fields_filled_%": chunk_fill,
                    "table_chunks_with_table_id_%": pct(sum(filled(r.get("table_id")) for r in tables), len(tables)),
                    **dict(c)}
    return res


# ---------- xlsx ----------

def load_xlsx_heuristics():
    """Эвристики XlsxParser проекта без src.config (заглушки), чтобы режим --view ничего не создавал.
    После переделки xlsx.py эвристик нет — тогда None."""
    try:
        for name in ("pcstub", "pcstub.parsers"):
            mod = types.ModuleType(name)
            mod.__path__ = []
            sys.modules.setdefault(name, mod)
        base = types.ModuleType("pcstub.parsers.base")
        base.Parser = type("Parser", (), {"__init__": lambda self: None})
        base.Chunk, base.parser = dict, (lambda *a: (lambda c: c))
        sys.modules["pcstub.parsers.base"] = base
        if "src.utils" not in sys.modules:
            import logging
            mod = types.ModuleType("src")
            mod.__path__ = []
            sys.modules["src"] = mod
            u = types.ModuleType("src.utils")
            u.__path__ = []
            u.get_logger = logging.getLogger
            sys.modules["src.utils"] = u
            r7 = types.ModuleType("src.utils.r7_office")
            r7.convert_via_r7 = r7.find_x2t = lambda *a, **k: None
            sys.modules["src.utils.r7_office"] = r7
        spec = importlib.util.spec_from_file_location("pcstub.parsers.xlsx", PROJECT / "src/parsers/xlsx.py")
        x = importlib.util.module_from_spec(spec)
        sys.modules["pcstub.parsers.xlsx"] = x
        spec.loader.exec_module(x)
        return x.XlsxParser()
    except Exception:
        return None


@functools.lru_cache(maxsize=16)
def xlsx_sheets(path) -> dict[str, list[tuple]]:
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        return {sn: list(wb[sn].iter_rows(values_only=True)) for sn in wb.sheetnames}
    finally:
        wb.close()


def _num(s):
    try:
        return float(s)
    except (TypeError, ValueError):
        return None


def derive_xlsx_gold():
    """Эталон ответа xlsx-запроса: лист(ы), строка, период, значение. Выводится из xlsx и
    work/querygen/xlsx_rows.json: запрос сгенерирован по полю `chunk` (чанк старого парсера);
    поля `label`/`row` там — другая строка с тем же началом подписи, поэтому подпись и значение
    берутся из чанка. Период — дата («2026-03-01 00:00:00» в чанке) или строка шапки («9м24»).
    Значение «x» (у шапки из трёх колонок Рубли/Валюта/Всего в рублях бывает «x») — первое
    число правее в той же группе колонок: в чанке это следующие строки «Column_N».
    Если подходят несколько листов, сохраняются все — проверка примет любой."""
    rows = json.load(open(WORK / "querygen/xlsx_rows.json", encoding="utf-8"))
    queries = {q["qid"]: q for q in jl(DATA / "queries.jsonl")}
    out, missed = [], []
    for i, r in enumerate(rows):
        q = queries.get(f"xl{i:03d}")
        if not q:
            continue
        lines = [s.partition(": ") for s in r["chunk"].split("\n")]
        label = next((v for _, _, v in lines if norm(v).startswith(norm(q["row_label"]))), None)
        is_date = bool(re.fullmatch(r"\d{4}-\d\d-\d\d", q["period"]))
        key = q["period"] + " 00:00:00" if is_date else q["period"]
        pos = next((j for j, (k, _, _) in enumerate(lines) if k == key), None)
        val = None
        if pos is not None:
            val = _num(lines[pos][2])
            for k, _, v in lines[pos + 1:]:
                if val is not None or not k.startswith("Column_"):
                    break
                val = _num(v)
        if label is None or val is None:
            missed.append((q["qid"], "подпись или значение не найдены в исходном чанке"))
            continue
        doc = next(x for x in MANIFEST if x["title"] + ".xlsx" == q["source_file"] or x["doc_id"].endswith(q["source_file"]))
        found = []
        for sn, srows in xlsx_sheets(doc["path"]).items():
            title = next((c for rr in srows[:6] for c in rr if isinstance(c, str) and re.match(r"\s*Таблица\s*\d", c)), None)
            groups = []  # колонки периода до следующей непустой ячейки той же строки шапки
            for rr in srows[:8]:
                for ci, c in enumerate(rr):
                    if (c.strftime("%Y-%m-%d") == q["period"] if isinstance(c, datetime.datetime)
                            else c is not None and norm(c) == norm(q["period"])):  # «9м24» или число 2025
                        end = next((cj for cj in range(ci + 1, len(rr)) if rr[cj] is not None), len(rr))
                        groups.append(range(ci, end))
            for row in srows:
                if any(isinstance(c, str) and norm(c) == norm(label) for c in row) and any(
                        ci < len(row) and isinstance(row[ci], (int, float)) and abs(row[ci] - val) < 1e-6 * max(1, abs(val))
                        for g in groups for ci in g):
                    found.append({"sheet": sn, "title": title})
        uniq = list({f["sheet"]: f for f in found}.values())
        if not uniq:
            missed.append((q["qid"], "строка не найдена в книге"))
            continue
        out.append({"qid": q["qid"], "doc_id": doc["doc_id"], "candidates": uniq,
                    "label": label, "period": q["period"], "value": val})
    return out, missed


def xlsx_gold():
    if XLSX_GOLD.exists():
        return jl(XLSX_GOLD)
    return derive_xlsx_gold()[0]


NUM_IN_TEXT = re.compile(r"-?\d[\d \u00a0\u202f]*(?:[.,]\d+)?")


def numbers_in(text: str) -> list[float]:
    out = []
    for s in NUM_IN_TEXT.findall(text):
        try:
            out.append(float(re.sub(r"[ \u00a0\u202f]", "", s).replace(",", ".")))
        except ValueError:
            pass
    return out


def check_xlsx(docs, show=("xl000", "xl003")):
    heur = load_xlsx_heuristics()
    files, stat = {}, collections.Counter()
    for d in docs.values():
        m = d["m"]
        if m["format"] != "xlsx":
            continue
        sheets = xlsx_sheets(m["path"])
        joined = norm("\n".join(d["chunks"]))
        # лист «попал в индекс», если в чанках есть его отличительная ячейка (есть только на этом листе)
        cells = {sn: {norm(c) for r in rows for c in r if isinstance(c, str) and len(c.strip()) >= 10} for sn, rows in sheets.items()}
        seen = collections.Counter(c for s in cells.values() for c in s)
        missing, unknown = [], []
        for sn, cs in cells.items():
            own = sorted(c for c in cs if seen[c] == 1 and not c.startswith("таблица"))[:5]
            if not own:
                unknown.append(sn)
            elif not any(c in joined for c in own):
                missing.append(sn)
        hdr = None
        if heur is not None and hasattr(heur, "_detect_headers_smart") and "1" in sheets:
            try:
                hdr = heur._detect_headers_smart(m["path"], "1")["header_row_index"]
            except Exception:
                hdr = None
        files[m["title"]] = {"sheets": len(sheets), "sheets_missing_in_chunks": missing, "sheets_no_marker": unknown,
                             "header_row_sheet_1": hdr}
        for x in d["chunks"]:
            lines = [s for s in x.split("\n") if s.strip()]
            stat["chunks"] += 1
            stat["lines"] += len(lines)
            stat["lines_Column_N"] += sum(s.startswith("Column_") for s in lines)
            stat["chunks_iso_datetime"] += bool(re.search(r"\d{4}-\d\d-\d\d 00:00:00", x))
            stat["chunks_long_float"] += bool(re.search(r"\d+\.\d{7,}", x))
            stat["chunks_human_date"] += bool(re.search(r"\b\d{1,2} (%s) \d{4}|\b\d\d\.\d\d\.\d{4}" % "|".join(MONTHS), x))
            stat["chunks_table_title"] += "Таблица" in x
    n = stat["chunks"]
    summary = {"chunks": n, "Column_N_lines_%": pct(stat["lines_Column_N"], stat["lines"]),
               **{k + "_%": pct(v, n) for k, v in stat.items() if k.startswith("chunks_")}}
    ans, shown = collections.Counter(), {}
    for g in xlsx_gold():
        human = None  # периоды-строки шапки («9м24», «2025») остаются как есть
        if re.fullmatch(r"\d{4}-\d\d-\d\d", g["period"]):
            y, mo, dd = map(int, g["period"].split("-"))
            human = [f"{dd} {MONTHS[mo - 1]} {y}", f"{dd:02d}.{mo:02d}.{y}"]
        titles = [c.get("title") or "" for c in g.get("candidates", [{"title": g.get("title")}])]
        title_w = [w[:6] for w in tok(re.sub(r"^\s*Таблица\s*[\d.]+", "", titles[0]))][:8]
        tol = max(0.051, 1e-6 * abs(g["value"]))
        hits = [x for x in docs[g["doc_id"]]["chunks"] if norm(g["label"]) in norm(x) and any(abs(v - g["value"]) <= tol for v in numbers_in(x))]
        ans["gold"] += 1
        if not hits:
            continue
        best = max(hits, key=lambda x: sum(w in norm(x) for w in title_w))
        ans["answer_chunk"] += 1
        ans["date_period"] += human is not None
        ans["human_date"] += human is not None and any(h in norm(best) for h in human)
        ans["title_50%"] += bool(title_w) and sum(w in norm(best) for w in title_w) / len(title_w) >= 0.5
        if g["qid"] in show:
            shown[g["qid"]] = best[:500]
    return {"summary": summary, "files": files,
            "answers": {"gold_queries": ans["gold"], "answer_chunk": ans["answer_chunk"],
                        "answer_chunk_date_period": ans["date_period"], "with_human_date": ans["human_date"],
                        "with_table_title_50%": ans["title_50%"]},
            "examples": shown}


# ---------- сравнение прогонов по срезам ----------

def sign_test_p(w: int, lose: int) -> float:
    n = w + lose
    if n == 0:
        return 1.0
    k = min(w, lose)
    return min(1.0, 2 * sum(math.comb(n, i) for i in range(k + 1)) / 2 ** n)


def query_slices(qids) -> dict[str, list[str]]:
    """Срезы запросов: итого, сценарий, формат и набор главного эталонного документа."""
    queries = {q["qid"]: q for q in jl(DATA / "queries.jsonl")}
    man = {m["doc_id"]: m for m in MANIFEST}
    gold = collections.defaultdict(dict)
    for line in open(DATA / "qrels.tsv", encoding="utf-8").read().splitlines()[1:]:
        q, d, g, _ = line.split("\t")
        gold[q][d] = int(g)
    slices = collections.defaultdict(list)
    for q in qids:
        mx = max(gold[q].values())
        prim = [d for d, g in gold[q].items() if g == mx]
        slices["Итого"].append(q)
        slices[f"сценарий {queries[q]['scenario']}"].append(q)
        for f in sorted({man[d]["format"] for d in prim}):
            slices[f"формат {f}"].append(q)
        for s in sorted({man[d]["dataset"] for d in prim}):
            slices[f"набор {s}"].append(q)
    return dict(sorted(slices.items(), key=lambda kv: (kv[0] != "Итого", kv[0])))


def compare_lines(run: dict, base: dict, qids, run_name: str, base_name: str,
                  metrics=("nDCG@10", "Hit@1", "Hit@10")) -> list[str]:
    """Парное сравнение двух выдач (qid → документы по порядку) по срезам: среднее, разница
    [95% ДИ, бутстрэп по кластерам], вариант «все неоценённые в топ-10 — правильные», знаки."""
    from metrics import bootstrap_ci, per_query
    from report import load_qrels
    queries = {q["qid"]: q for q in jl(DATA / "queries.jsonl")}
    qrels = load_qrels()
    qids = sorted(qids)
    pr, pb = per_query(run, qrels, qids), per_query(base, qrels, qids)
    # Крайние варианты учёта неоценённых: «все новые правильные» — оценка 2 каждому неоценённому.
    up = {q: {**qrels.get(q, {}), **{d: 2 for d in run.get(q, [])[:10] if d not in qrels.get(q, {})}} for q in qids}
    pr_up = per_query(run, up, qids)
    clusters = {q: queries[q]["base"] for q in queries}
    out = [f"| Срез | Запросов | метрика | {base_name} | {run_name} | разница [95% ДИ] | если новые правильные | "
           "лучше/хуже | p (знаки) | неоценённых в топ-10 |",
           "|---|---:|---|---:|---:|---:|---:|---:|---:|---:|"]
    for name, qs in query_slices(qids).items():
        unj = sum(1 for q in qs for d in run.get(q, [])[:10] if d not in qrels.get(q, {}))
        for met in metrics:
            diff = {q: {met: pr[q][met] - pb[q][met]} for q in qs}
            pt, lo, hi = bootstrap_ci(diff, met, clusters)
            up_mean = statistics.mean(pr_up[q][met] for q in qs) - statistics.mean(pb[q][met] for q in qs)
            w = sum(diff[q][met] > 1e-9 for q in qs)
            lose = sum(diff[q][met] < -1e-9 for q in qs)
            out.append(f"| {name} | {len(qs)} | {met} | {statistics.mean(pb[q][met] for q in qs):.3f} | "
                       f"{statistics.mean(pr[q][met] for q in qs):.3f} | {pt:+.3f} [{lo:+.3f}…{hi:+.3f}] | {up_mean:+.3f} | "
                       f"{w}/{lose} | {sign_test_p(w, lose):.3f} | {unj} |")
    return out


def cmd_compare(a):
    from metrics import relevant
    from report import corpus_rel, load_qrels, load_run
    queries = {q["qid"]: q for q in jl(DATA / "queries.jsonl")}
    qrels = load_qrels()

    def answerable(run):
        st = json.load(open(EVAL / "runs" / run / "index_state.json", encoding="utf-8"))
        idx = {corpus_rel(p) for p, n in st["probe_paths"].get("paths", {}).items() if n > 0}
        return {q for q in queries if queries[q]["scenario"] != 6 and relevant(qrels.get(q, {})) & idx}

    ans_run, ans_base = answerable(a.run), answerable(a.base)
    if ans_run != ans_base:
        print(f"ВНИМАНИЕ: разный набор запросов с эталоном в индексе: {a.run} {len(ans_run)}, {a.base} {len(ans_base)}; "
              f"сравнение по пересечению, выпали: {sorted(ans_base ^ ans_run)[:10]}")
    run, _ = load_run(EVAL / "runs" / a.run, "hybrid")
    base, _ = load_run(EVAL / "runs" / a.base, "hybrid")
    print("\n".join(compare_lines(run, base, ans_run & ans_base, a.run, a.base, a.metrics.split(","))))
    upper, _ = load_run(EVAL / "runs" / a.run, "hybrid_upper")
    if upper:  # прогон с --upper: те же запросы ЗАГЛАВНЫМИ на том же индексе
        same = sum(upper.get(q, [])[:10] == run.get(q, [])[:10] for q in run)
        print(f"\nРегистр ({a.run}): топ-10 запросов ЗАГЛАВНЫМИ совпал с обычными в {same} из {len(run)}\n")
        print("\n".join(compare_lines(upper, run, ans_run, f"{a.run} ЗАГЛАВНЫМИ", a.run, a.metrics.split(","))[:5]))


# ---------- команды ----------

def cmd_check(a):
    src = LiveSource() if a.live else ViewSource()
    t = time.perf_counter()
    docs = load_all(src)
    count = token_counter(src)
    rep = {"source": "live" if a.live else "view", "project": str(PROJECT) if a.live else None,
           "parser_keys": getattr(src, "keys", {}),
           "same_as_R7": check_same_as_view(docs) if a.live else None,
           "volume": check_volume(docs, count), "pdf_text": check_pdf_text(docs), "context": check_context(docs),
           "pptx": check_pptx(docs), "source_md_coverage": check_source_md(docs), "answers": check_answers(docs),
           "xlsx": check_xlsx(docs), "hygiene": check_hygiene(docs), "short_dups": check_short_dups(docs),
           "metadata": check_metadata(docs), "seconds": round(time.perf_counter() - t, 1)}
    text = json.dumps(rep, ensure_ascii=False, indent=1, default=str)
    print(text)
    if a.label:
        OUT.mkdir(parents=True, exist_ok=True)
        (OUT / f"{a.label}.json").write_text(text, encoding="utf-8")


def flatten(d, prefix=""):
    out = {}
    for k, v in d.items():
        if isinstance(v, dict):
            out.update(flatten(v, f"{prefix}{k}."))
        elif isinstance(v, (int, float)) and not isinstance(v, bool):
            out[prefix + k] = v
    return out


def cmd_delta(a):
    x, y = (flatten(json.load(open(p, encoding="utf-8"))) for p in (a.a, a.b))
    for k in sorted(set(x) | set(y)):
        if x.get(k) != y.get(k) and not k.endswith("parse_ms") and k != "seconds":
            print(f"{k}: {x.get(k)} → {y.get(k)}")


def cmd_freeze_xlsx_gold(a):
    if XLSX_GOLD.exists() and not a.force:
        raise SystemExit(f"{XLSX_GOLD} уже есть; эталон замораживается один раз (--force перезапишет)")
    gold, missed = derive_xlsx_gold()
    XLSX_GOLD.write_text("".join(json.dumps(g, ensure_ascii=False) + "\n" for g in gold), encoding="utf-8")
    amb = sum(len(g["candidates"]) > 1 for g in gold)
    print(f"записано {len(gold)} эталонов xlsx (из них с несколькими подходящими листами: {amb}); не выведено: {len(missed)}")
    for qid, why in missed:
        print(f"  {qid}: {why}")


def main():
    ap = argparse.ArgumentParser()
    sub = ap.add_subparsers(dest="cmd", required=True)
    c = sub.add_parser("check")
    g = c.add_mutually_exclusive_group(required=True)
    g.add_argument("--view", action="store_true", help="чанки из work/indexer_view_corpus (базовая линия R7)")
    g.add_argument("--live", action="store_true", help="чанки текущими парсерами проекта (с кэшем)")
    c.add_argument("--label", help="сохранить отчёт в work/parse_checks/<label>.json")
    d = sub.add_parser("delta")
    d.add_argument("a")
    d.add_argument("b")
    r = sub.add_parser("compare")
    r.add_argument("--run", required=True)
    r.add_argument("--base", default="R7")
    r.add_argument("--metrics", default="nDCG@10,Hit@1,Hit@10")
    f = sub.add_parser("freeze-xlsx-gold")
    f.add_argument("--force", action="store_true")
    a = ap.parse_args()
    {"check": cmd_check, "delta": cmd_delta, "compare": cmd_compare, "freeze-xlsx-gold": cmd_freeze_xlsx_gold}[a.cmd](a)


if __name__ == "__main__":
    main()
